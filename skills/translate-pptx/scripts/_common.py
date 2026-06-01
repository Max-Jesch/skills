"""Shared helpers for the translate-pptx skill scripts.

The single source of truth here is `iter_translatable_frames`. Both extract.py
and apply.py walk the deck through this generator, so a frame's *position* in
the walk is its stable `id`. Never change the walk order without bumping both
scripts together, or extracted translations will map back to the wrong frames.
"""
from __future__ import annotations

from pptx.enum.shapes import MSO_SHAPE_TYPE


def iter_translatable_frames(prs, include_notes=True):
    """Yield a dict per translatable TextFrame, in a deterministic order.

    Each item: {"kind", "slide", "autofit_ok", "tf"}.
      kind         - "slide" | "table" | "chart_title" | "notes"
      slide        - 1-based slide number (informational)
      autofit_ok   - whether it is safe to apply shrink-to-fit autofit here
      tf           - the python-pptx TextFrame object
    """
    for s_idx, slide in enumerate(prs.slides):
        yield from _walk(slide.shapes, s_idx + 1)
        if include_notes and slide.has_notes_slide:
            tf = slide.notes_slide.notes_text_frame
            if tf is not None:
                yield {"kind": "notes", "slide": s_idx + 1,
                       "autofit_ok": False, "tf": tf}


def _walk(shapes, slide_no):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shape.shapes, slide_no)
            continue
        if shape.has_text_frame:
            yield {"kind": "slide", "slide": slide_no,
                   "autofit_ok": True, "tf": shape.text_frame}
        elif getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield {"kind": "table", "slide": slide_no,
                           "autofit_ok": False, "tf": cell.text_frame}
        elif getattr(shape, "has_chart", False) and shape.has_chart:
            try:
                chart = shape.chart
                if chart.has_title:
                    yield {"kind": "chart_title", "slide": slide_no,
                           "autofit_ok": False,
                           "tf": chart.chart_title.text_frame}
            except Exception:
                pass
