"""Shared helpers for the edit-pptx skill scripts.

Two jobs live here:

1. **Addressing** — `walk_shapes` flattens groups into a deterministic order and
   every shape is keyed by its slide number + python-pptx `shape_id` (the
   `<p:cNvPr id>`, unique within a slide and *stable* when other shapes are
   added or removed). `inspect.py` reports those ids; `apply_ops.py` resolves
   them with `find_shape`. Positional walk indices are deliberately NOT used as
   keys, because inserting/deleting shapes shifts them.

2. **Fit-aware text writing** — the same machinery translate-pptx uses: write
   into the first run to keep typography, re-materialise line breaks as real
   <a:br/>, turn on shrink-to-fit autofit, and (optionally) scale the real run
   sizes so LibreOffice/PDF render the same fit PowerPoint computes on open.
"""
from __future__ import annotations

import copy

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn

EMU_PER_INCH = 914400

# python-pptx surfaces an in-paragraph break (<a:br/>) as a vertical tab in
# .text; a literal newline can appear the same way. Assigning either straight to
# run.text would escape it to the visible literal "_x000B_", so we split on these
# and rebuild real <a:br/> elements instead.
_LINE_BREAK_CHARS = ("\x0b", "\n", "\r")


# --------------------------------------------------------------------------- #
# Addressing
# --------------------------------------------------------------------------- #
def walk_shapes(shapes, _depth=0):
    """Yield (shape, depth) for every shape, descending into groups.

    depth 0 is a top-level shape on the slide; depth > 0 means it lives inside a
    group. The group container itself is yielded too (before its children).
    """
    for shape in shapes:
        yield shape, _depth
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk_shapes(shape.shapes, _depth + 1)


def find_shape(slide, shape_id):
    """Return the shape with this shape_id on `slide` (searching groups), or None."""
    for shape, _ in walk_shapes(slide.shapes):
        if shape.shape_id == shape_id:
            return shape
    return None


def autofit_ok(shape):
    """Whether it is safe to apply slide shrink-to-fit autofit to this shape."""
    return shape.has_text_frame and shape.shape_type != MSO_SHAPE_TYPE.GROUP


# --------------------------------------------------------------------------- #
# Text writing (format-preserving, line-break-safe)
# --------------------------------------------------------------------------- #
def _split_line_breaks(text):
    segments = [text]
    for ch in _LINE_BREAK_CHARS:
        segments = [piece for seg in segments for piece in seg.split(ch)]
    return segments


def write_paragraph_text(para, text, warn=None):
    """Set `para`'s text, preserving the first run's formatting.

    Line breaks in `text` become real <a:br/> elements (cloned from the first
    run) rather than an escaped "_x000B_" string. An empty paragraph gets a
    single default-formatted run.
    """
    segments = _split_line_breaks(text)
    runs = para.runs
    if runs:
        first = runs[0]
        first.text = segments[0]
        for r in runs[1:]:
            r.text = ""
    else:
        para.text = segments[0]  # creates one run with default formatting
        first = para.runs[0]

    if len(segments) == 1:
        return

    r_el = first._r
    rPr = r_el.find(qn("a:rPr"))
    parent = r_el.getparent()
    insert_at = list(parent).index(r_el) + 1
    for seg in segments[1:]:
        br = r_el.makeelement(qn("a:br"), {})
        if rPr is not None:
            br.append(copy.deepcopy(rPr))
        new_r = copy.deepcopy(r_el)
        new_t = new_r.find(qn("a:t"))
        if new_t is None:
            new_t = new_r.makeelement(qn("a:t"), {})
            new_r.append(new_t)
        new_t.text = seg
        parent.insert(insert_at, br)
        parent.insert(insert_at + 1, new_r)
        insert_at += 2


def add_paragraph_like(tf, text, level=None, warn=None):
    """Append a paragraph to `tf`, inheriting the last paragraph's bullet/format.

    The last existing <a:p> is deep-copied (so its bullet, indent, level and run
    properties carry over), stripped down to a single run, then re-texted. Falls
    back to a plain default paragraph when the frame is empty.
    """
    paras = tf.paragraphs
    if paras and paras[-1].runs:
        template = paras[-1]._p
        new_p = copy.deepcopy(template)
        for br in new_p.findall(qn("a:br")):
            new_p.remove(br)
        runs_el = new_p.findall(qn("a:r"))
        for extra in runs_el[1:]:
            new_p.remove(extra)
        template.getparent().append(new_p)
        new_para = tf.paragraphs[-1]
    else:
        new_para = tf.add_paragraph()
    write_paragraph_text(new_para, text, warn)
    if level is not None:
        new_para.level = level
    return new_para


def delete_paragraph(tf, p_idx, warn=None):
    paras = tf.paragraphs
    if p_idx >= len(paras):
        if warn:
            warn(f"paragraph index {p_idx} out of range")
        return
    p_el = paras[p_idx]._p
    p_el.getparent().remove(p_el)


def set_frame_text(tf, text, warn=None):
    """Set the whole frame's text. `text` is a str (one paragraph) or a list of
    paragraph strings; surplus existing paragraphs are removed."""
    lines = text if isinstance(text, list) else [text]
    paras = tf.paragraphs
    for i, line in enumerate(lines):
        if i < len(paras):
            write_paragraph_text(paras[i], line, warn)
        else:
            add_paragraph_like(tf, line, warn=warn)
    for p in list(tf.paragraphs)[len(lines):]:
        p._p.getparent().remove(p._p)


# --------------------------------------------------------------------------- #
# Fit (autofit + real-size scaling) — see reference.md for the why
# --------------------------------------------------------------------------- #
def _scale_explicit_sizes(tf, factor):
    """Multiply every explicit font size (sz, 1/100 pt) in the frame by factor.

    LibreOffice and PDF export ignore normAutofit fontScale, so scaling the real
    run sizes is what makes a forced shrink render the same everywhere. Runs that
    inherit their size from the placeholder/theme have no sz and are left to
    normAutofit (PowerPoint re-fits them on open)."""
    txBody = tf._txBody
    scaled = 0
    for tag in ("a:rPr", "a:defRPr", "a:endParaRPr"):
        for el in txBody.iter(qn(tag)):
            sz = el.get("sz")
            if sz is None:
                continue
            el.set("sz", str(max(100, int(round(int(sz) * factor)))))
            scaled += 1
    return scaled


def set_autofit(tf, font_scale=None, line_reduction=None):
    """Enable shrink-to-fit autofit; optionally force a real-size scale.

    A bare <a:normAutofit/> lets PowerPoint re-fit on open. When font_scale is
    given we ALSO shrink the explicit run sizes (see _scale_explicit_sizes) so
    every renderer honours it; we deliberately do not bake fontScale too, to
    avoid PowerPoint double-shrinking."""
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE  # writes a well-formed normAutofit
    if font_scale is not None:
        _scale_explicit_sizes(tf, float(font_scale) / 100.0)
    if line_reduction is not None:
        bodyPr = tf._txBody.find(qn("a:bodyPr"))
        na = bodyPr.find(qn("a:normAutofit")) if bodyPr is not None else None
        if na is not None:
            na.set("lnSpcReduction", str(int(round(float(line_reduction) * 1000))))
