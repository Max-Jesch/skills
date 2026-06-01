#!/usr/bin/env python3
"""Extract translatable text from a .pptx into a flat JSON job file.

Usage:
    python extract.py deck.pptx --target de -o job.json

The JSON it writes is the thing *you* (the model) translate: fill each
paragraph's empty "tgt" with a concise translation, leave "id"/"i"/"src"
untouched, then hand the file to apply.py.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _common import iter_translatable_frames  # noqa: E402


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("pptx", help="source presentation")
    ap.add_argument("--target", default="de",
                    help="target language code (default: de)")
    ap.add_argument("--no-notes", action="store_true",
                    help="skip speaker notes")
    ap.add_argument("-o", "--out", default="-",
                    help="output JSON path (default: stdout)")
    args = ap.parse_args()

    include_notes = not args.no_notes
    prs = Presentation(args.pptx)

    frames = []
    for fid, item in enumerate(
            iter_translatable_frames(prs, include_notes=include_notes)):
        paras = []
        for p_idx, para in enumerate(item["tf"].paragraphs):
            if para.text.strip():
                paras.append({"i": p_idx, "src": para.text, "tgt": ""})
        frames.append({
            "id": fid,
            "slide": item["slide"],
            "kind": item["kind"],
            "autofit": item["autofit_ok"],
            "paras": paras,
        })

    doc = {
        "source": str(Path(args.pptx).name),
        "target_lang": args.target,
        "include_notes": include_notes,
        "frames": frames,
    }
    text = json.dumps(doc, ensure_ascii=False, indent=2)
    if args.out == "-":
        sys.stdout.write(text + "\n")
    else:
        Path(args.out).write_text(text + "\n", encoding="utf-8")

    n_para = sum(len(f["paras"]) for f in frames)
    print(f"extracted {len(frames)} text frames / {n_para} paragraphs "
          f"across {len(prs.slides)} slides", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
