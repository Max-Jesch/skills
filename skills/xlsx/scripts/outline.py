#!/usr/bin/env python3
"""Print a compact map of an .xlsx so you can orient before editing it.

Dumps per-sheet dimensions, the header row, a few sample rows, merged ranges,
and workbook-level defined names — without slurping the whole file into context.
Use this first on any existing workbook: it tells you the real column layout and
where data starts, so your formulas reference the right cells.

By default it shows formulas as written. Pass --values to show cached results
instead (run recalc.py first, or cached values will be blank for cells openpyxl
wrote).
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter


def fmt(v) -> str:
    if v is None:
        return ""
    s = str(v)
    return s if len(s) <= 28 else s[:25] + "..."


def main() -> int:
    ap = argparse.ArgumentParser(description="Summarize the structure of an .xlsx.")
    ap.add_argument("file", help="path to the .xlsx")
    ap.add_argument("--rows", type=int, default=5, help="sample data rows per sheet (default 5)")
    ap.add_argument("--values", action="store_true", help="show cached values instead of formulas")
    ap.add_argument("--sheet", help="only inspect this sheet")
    args = ap.parse_args()

    path = Path(args.file)
    if not path.exists():
        print(f"error: {path} not found", file=sys.stderr)
        return 2

    wb = openpyxl.load_workbook(path, data_only=args.values)
    names = list(wb.defined_names) if hasattr(wb.defined_names, "__iter__") else []

    print(f"# {path.name}")
    print(f"sheets: {', '.join(wb.sheetnames)}")
    if names:
        print(f"defined names: {', '.join(names)}")
    print()

    targets = [args.sheet] if args.sheet else wb.sheetnames
    for name in targets:
        ws = wb[name]
        print(f"## {name}  ({ws.max_row} rows x {ws.max_column} cols, dims {ws.dimensions})")
        if ws.merged_cells.ranges:
            mr = ", ".join(str(r) for r in list(ws.merged_cells.ranges)[:10])
            extra = "" if len(ws.merged_cells.ranges) <= 10 else " ..."
            print(f"merged: {mr}{extra}")

        ncols = min(ws.max_column, 12)
        header = [get_column_letter(c) for c in range(1, ncols + 1)]
        print("col:   " + " | ".join(f"{h:>10}" for h in header))
        nrows = min(ws.max_row, args.rows + 1)
        for r in range(1, nrows + 1):
            cells = [fmt(ws.cell(row=r, column=c).value) for c in range(1, ncols + 1)]
            print(f"r{r:<4} " + " | ".join(f"{c:>10}" for c in cells))
        if ws.max_column > ncols:
            print(f"   (+{ws.max_column - ncols} more columns)")
        print()

    wb.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
