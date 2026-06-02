#!/usr/bin/env python3
"""Inventory a .pptx into a structured JSON map you read before editing.

Usage:
    python3 inspect.py deck.pptx -o deck.json

Unlike a translation extract (which only needs text), this reports EVERY shape
with the address you edit by — slide number + shape_id — plus its type, geometry
(in inches) and text, the table grid where present, and the deck's available
slide layouts (so you know what add_slide / duplicate_slide can target).

Read the JSON together with the rendered slides: the render tells you what the
deck looks like, this tells you how to address what you want to change.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _common import EMU_PER_INCH, walk_shapes  # noqa: E402


def _in(v):
    return round(v / EMU_PER_INCH, 2) if v is not None else None


def _type(shape):
    st = shape.shape_type
    if st is None:
        return "UNKNOWN"
    return str(st).split(" (")[0]


def _shape_record(shape, depth):
    rec = {
        "shape_id": shape.shape_id,
        "name": shape.name,
        "type": _type(shape),
        "left_in": _in(shape.left),
        "top_in": _in(shape.top),
        "width_in": _in(shape.width),
        "height_in": _in(shape.height),
    }
    if depth > 0:
        rec["in_group"] = True
    if shape.has_text_frame:
        paras = [{"p": i, "text": p.text}
                 for i, p in enumerate(shape.text_frame.paragraphs)
                 if p.text.strip()]
        if paras:
            rec["paras"] = paras
    if getattr(shape, "has_table", False) and shape.has_table:
        t = shape.table
        ncols = len(t.columns)
        rec["table"] = {
            "rows": len(t.rows),
            "cols": ncols,
            "cells": [[t.cell(r, c).text for c in range(ncols)]
                      for r in range(len(t.rows))],
        }
    if getattr(shape, "has_chart", False) and shape.has_chart:
        rec["chart"] = True
    return rec


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx", help="presentation to inspect")
    ap.add_argument("-o", "--out", default="-", help="output JSON (default: stdout)")
    args = ap.parse_args()

    prs = Presentation(args.pptx)

    slides = []
    for s_idx, slide in enumerate(prs.slides):
        shapes = [_shape_record(sh, d) for sh, d in walk_shapes(slide.shapes)]
        slides.append({
            "slide": s_idx + 1,
            "layout": slide.slide_layout.name,
            "shapes": shapes,
        })

    doc = {
        "source": Path(args.pptx).name,
        "slide_count": len(prs.slides),
        "slide_size": {"width_in": _in(prs.slide_width),
                       "height_in": _in(prs.slide_height)},
        "layouts": [{"index": i, "name": l.name}
                    for i, l in enumerate(prs.slide_layouts)],
        "slides": slides,
    }

    text = json.dumps(doc, ensure_ascii=False, indent=2)
    if args.out == "-":
        sys.stdout.write(text + "\n")
    else:
        Path(args.out).write_text(text + "\n", encoding="utf-8")

    n_shapes = sum(len(s["shapes"]) for s in slides)
    print(f"inspected {len(slides)} slides / {n_shapes} shapes; "
          f"{len(doc['layouts'])} layouts available", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
