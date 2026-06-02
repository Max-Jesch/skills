#!/usr/bin/env python3
"""Execute a declarative ops list against a .pptx and write a new file.

Usage:
    python3 apply_ops.py deck.pptx ops.json -o deck.edited.pptx

You author a small `ops.json` (a list of operations); this runs them in order,
preserving formatting and turning on shrink-to-fit autofit on any slide text box
it touches. Never hand-author a giant deck dump — keep the ops file small and
let inspect.py tell you the addresses.

Addresses are slide number (1-based) + shape_id (from inspect.py). For anything
the op vocabulary below does not cover (restyling, move/resize, deleting shapes,
exotic objects) reach for the python-pptx escape hatch in reference.md.

Ops (one object per entry in the top-level "ops" list):

  TEXT
    {"op":"set_text","slide":2,"shape_id":5,"text":"New title"}
    {"op":"set_text","slide":2,"shape_id":5,"text":["Bullet A","Bullet B"]}
    {"op":"set_text","slide":2,"shape_id":5,"p":0,"text":"Just paragraph 0"}
    {"op":"set_text","slide":2,"shape_id":9,"row":1,"col":0,"text":"cell"}   # table cell
    {"op":"add_bullet","slide":2,"shape_id":6,"text":"Another point","level":1}
    {"op":"delete_paragraph","slide":2,"shape_id":6,"p":3}
       (set_text on a slide text box also accepts "font_scale" 1-100 and
        "line_reduction" 0-20 to force a tighter fit — pick by eye from renders.)

  SLIDES
    {"op":"duplicate_slide","from":3,"at":4,"set":{"5":"Title","6":["a","b"]}}
       (deep-copies slide 3 -> new slide at position 4; "set" fills text in the
        copy keyed by the SOURCE shape_id, optionally "<id>.<p>" for one para.)
    {"op":"add_slide","layout":1,"at":5}
    {"op":"delete_slide","slide":7}
    {"op":"move_slide","slide":8,"to":2}

  MEDIA (positions in inches)
    {"op":"add_textbox","slide":4,"text":"Caption","left":1,"top":6,"width":8,"height":0.8}
    {"op":"add_picture","slide":4,"path":"logo.png","left":0.5,"top":0.3,"width":1.5}
    {"op":"add_table","slide":4,"rows":[["A","B"],["1","2"]],"left":1,"top":2,"width":6,"height":3}

Structural slide ops shift slide numbers. If you add/delete/move a slide and then
need to edit it by number, run apply_ops, then RE-INSPECT, then a second ops file
— or use duplicate_slide's inline "set".
"""
from __future__ import annotations

import argparse
import copy
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _common import (  # noqa: E402
    add_paragraph_like, autofit_ok, delete_paragraph, find_shape,
    set_autofit, set_frame_text, write_paragraph_text)

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


# --------------------------------------------------------------------------- #
# Slide-level structural helpers (python-pptx has no public API for these)
# --------------------------------------------------------------------------- #
def _remap_rids(el, rid_map):
    """Rewrite relationship-id attributes (r:embed, r:id, r:link) in a copied
    subtree to the ids they were re-related to on the destination slide."""
    prefix = "{%s}" % R_NS
    for node in el.iter():
        for name, val in list(node.attrib.items()):
            if name.startswith(prefix) and val in rid_map:
                node.set(name, rid_map[val])


def _duplicate_slide(prs, src_index):
    """Deep-copy the slide at src_index, append it, and return the new slide.

    Copies the source's relationships (images, hyperlinks, media) to fresh ids
    and remaps the references in the cloned shape XML. Handles typical decks
    (text + pictures + hyperlinks); very complex objects (charts with embedded
    workbooks, SmartArt, media) may need manual duplication in PowerPoint."""
    source = prs.slides[src_index]
    dest = prs.slides.add_slide(source.slide_layout)

    # Drop the placeholders the layout injected — we're replacing the body wholesale.
    for shp in list(dest.shapes):
        shp._element.getparent().remove(shp._element)

    rid_map = {}
    for rId, rel in source.part.rels.items():
        if rel.reltype.endswith("slideLayout"):
            continue
        new_rid = dest.part.relate_to(rel._target, rel.reltype,
                                      is_external=rel.is_external)
        rid_map[rId] = new_rid

    dest_tree = dest.shapes._spTree
    for child in source.shapes._spTree:
        tag = child.tag
        if tag.endswith("}nvGrpSpPr") or tag.endswith("}grpSpPr"):
            continue
        new_child = copy.deepcopy(child)
        _remap_rids(new_child, rid_map)
        dest_tree.append(new_child)
    return dest


def _move_slide(prs, from_index, to_index):
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    el = ids[from_index]
    sldIdLst.remove(el)
    sldIdLst.insert(to_index, el)


def _delete_slide(prs, index):
    sldIdLst = prs.slides._sldIdLst
    sldId = list(sldIdLst)[index]
    prs.part.drop_rel(sldId.get(qn("r:id")))
    sldIdLst.remove(sldId)


# --------------------------------------------------------------------------- #
# Op handlers — each takes (prs, op, warn)
# --------------------------------------------------------------------------- #
def _resolve_text_shape(prs, op, warn):
    slide = prs.slides[op["slide"] - 1]
    shape = find_shape(slide, op["shape_id"])
    if shape is None:
        warn(f"shape_id {op['shape_id']} not found on slide {op['slide']}")
    return shape


def op_set_text(prs, op, warn):
    shape = _resolve_text_shape(prs, op, warn)
    if shape is None:
        return
    text = op["text"]

    if "row" in op and "col" in op:
        if not (getattr(shape, "has_table", False) and shape.has_table):
            warn(f"shape_id {op['shape_id']} is not a table")
            return
        set_frame_text(shape.table.cell(op["row"], op["col"]).text_frame, text, warn)
        return

    if not shape.has_text_frame:
        warn(f"shape_id {op['shape_id']} has no text frame")
        return
    tf = shape.text_frame
    if "p" in op:
        if op["p"] >= len(tf.paragraphs):
            warn(f"paragraph {op['p']} out of range on shape {op['shape_id']}")
            return
        write_paragraph_text(tf.paragraphs[op["p"]], text, warn)
    else:
        set_frame_text(tf, text, warn)

    if autofit_ok(shape):
        set_autofit(tf, op.get("font_scale"), op.get("line_reduction"))


def op_add_bullet(prs, op, warn):
    shape = _resolve_text_shape(prs, op, warn)
    if shape is None:
        return
    if not shape.has_text_frame:
        warn(f"shape_id {op['shape_id']} has no text frame")
        return
    add_paragraph_like(shape.text_frame, op["text"], op.get("level"), warn)
    if autofit_ok(shape):
        set_autofit(shape.text_frame)


def op_delete_paragraph(prs, op, warn):
    shape = _resolve_text_shape(prs, op, warn)
    if shape is None:
        return
    if not shape.has_text_frame:
        warn(f"shape_id {op['shape_id']} has no text frame")
        return
    delete_paragraph(shape.text_frame, op["p"], warn)
    if autofit_ok(shape):
        set_autofit(shape.text_frame)


def _apply_set_map(slide, mapping, warn):
    for key, val in mapping.items():
        if "." in key:
            sid, p = key.split(".", 1)
            sid, p = int(sid), int(p)
        else:
            sid, p = int(key), None
        shape = find_shape(slide, sid)
        if shape is None or not shape.has_text_frame:
            warn(f"set: shape_id {sid} not a text shape on duplicated slide")
            continue
        tf = shape.text_frame
        if p is None:
            set_frame_text(tf, val, warn)
        else:
            if p >= len(tf.paragraphs):
                warn(f"set: paragraph {p} out of range on shape {sid}")
                continue
            write_paragraph_text(tf.paragraphs[p], val, warn)
        if autofit_ok(shape):
            set_autofit(tf)


def op_duplicate_slide(prs, op, warn):
    _duplicate_slide(prs, op["from"] - 1)
    new_index = len(prs.slides._sldIdLst) - 1
    if "at" in op:
        _move_slide(prs, new_index, op["at"] - 1)
        new_index = op["at"] - 1
    if "set" in op:
        _apply_set_map(prs.slides[new_index], op["set"], warn)


def op_add_slide(prs, op, warn):
    prs.slides.add_slide(prs.slide_layouts[op["layout"]])
    if "at" in op:
        _move_slide(prs, len(prs.slides._sldIdLst) - 1, op["at"] - 1)


def op_delete_slide(prs, op, warn):
    _delete_slide(prs, op["slide"] - 1)


def op_move_slide(prs, op, warn):
    _move_slide(prs, op["slide"] - 1, op["to"] - 1)


def op_add_textbox(prs, op, warn):
    slide = prs.slides[op["slide"] - 1]
    tb = slide.shapes.add_textbox(Inches(op["left"]), Inches(op["top"]),
                                  Inches(op.get("width", 4)),
                                  Inches(op.get("height", 1)))
    set_frame_text(tb.text_frame, op["text"], warn)
    if op.get("font_scale") is not None or op.get("line_reduction") is not None:
        set_autofit(tb.text_frame, op.get("font_scale"), op.get("line_reduction"))


def op_add_picture(prs, op, warn):
    slide = prs.slides[op["slide"] - 1]
    w = Inches(op["width"]) if "width" in op else None
    h = Inches(op["height"]) if "height" in op else None
    slide.shapes.add_picture(op["path"], Inches(op["left"]), Inches(op["top"]), w, h)


def op_add_table(prs, op, warn):
    slide = prs.slides[op["slide"] - 1]
    rows = op["rows"]
    nrows, ncols = len(rows), max(len(r) for r in rows)
    table = slide.shapes.add_table(
        nrows, ncols, Inches(op["left"]), Inches(op["top"]),
        Inches(op.get("width", 6)), Inches(op.get("height", nrows * 0.4))).table
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.cell(ri, ci).text = str(val)


HANDLERS = {
    "set_text": op_set_text,
    "add_bullet": op_add_bullet,
    "delete_paragraph": op_delete_paragraph,
    "duplicate_slide": op_duplicate_slide,
    "add_slide": op_add_slide,
    "delete_slide": op_delete_slide,
    "move_slide": op_move_slide,
    "add_textbox": op_add_textbox,
    "add_picture": op_add_picture,
    "add_table": op_add_table,
}


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx", help="source presentation")
    ap.add_argument("ops", help="ops JSON")
    ap.add_argument("-o", "--out", help="output .pptx (default: <stem>.edited.pptx)")
    args = ap.parse_args()

    doc = json.loads(Path(args.ops).read_text(encoding="utf-8"))
    ops = doc["ops"] if isinstance(doc, dict) else doc

    prs = Presentation(args.pptx)
    warnings: list[str] = []
    done = 0
    for n, op in enumerate(ops):
        kind = op.get("op")
        handler = HANDLERS.get(kind)
        if handler is None:
            warnings.append(f"op {n}: unknown op '{kind}'")
            continue

        def warn(msg, _n=n, _k=kind):
            warnings.append(f"op {_n} ({_k}): {msg}")

        try:
            handler(prs, op, warn)
            done += 1
        except Exception as e:  # one bad op shouldn't sink the rest
            warnings.append(f"op {n} ({kind}): FAILED — {type(e).__name__}: {e}")

    out = args.out or f"{Path(args.pptx).stem}.edited.pptx"
    prs.save(out)

    print(f"applied {done}/{len(ops)} ops -> {out}", file=sys.stderr)
    if warnings:
        print(f"{len(warnings)} warning(s):", file=sys.stderr)
        for w in warnings[:50]:
            print(f"  - {w}", file=sys.stderr)
        if len(warnings) > 50:
            print(f"  ... and {len(warnings) - 50} more", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
